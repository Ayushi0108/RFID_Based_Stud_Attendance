from django.shortcuts import render, redirect
# from django.http import HttpResponse
from math import ceil
from openpyxl import load_workbook
from user_authentication.models import Stud   
import datetime as dt
from user_authentication.models import Record
# Create your views here.

def dashboard(request):
    print("start")
    if not request.user.is_authenticated:
        return redirect('login')
    student = Stud.objects.filter(email=request.user.username).first()
    username = student.name
    roll_number = student.roll
    
    course_info=Record.objects.all().order_by('course_code').values('course_code','course')
    courses=[]
    record=[]
    for course in course_info:
        # print(course)
        course_code = course['course_code']
        course_name = course['course']
        # Creating courses string to display Registered courses
        string = course_code +" - "+ course_name
        courses.append(string)
        record.append(get_attendance_info(course_code,course_name, roll_number))

    today_date = dt.datetime.now()
    today_date = f'{today_date:%d-%m-%Y}'
    print("end")
    return render(request,'student_dashboard/index.html',
    {
    'record':record,
    'username':username,
    'roll_number':roll_number,
    'courses':courses,
    'today_date':today_date,
    # 'graph_data':[1,2,3]
    })


# Functions for dashboard
excel_sheet_path='Excel_Sheets/main/'

def get_attendance_info(course_code, course_name, roll_number):
    # return {'cid':course_code, 'cname':course_name,'stot_lect': 2,'sper':'78'}
    sheet_name = course_code+'.xlsx'
    # Loading the excel sheet
    try:
        wb = load_workbook(excel_sheet_path+sheet_name)
    except Exception:
        return {'cid':course_code, 'cname':course_name,'stot_lect': '-','total_lect':'-','sper': '-', 'detained_status':'yes'}

    sheet = wb.active
    # sheet[index] - selects a particular row in the excel sheet
    # in our case as the roll number is a simple digit like 1,2,3 and not something like 21105002, etc
    # it becomes easy to quickly fetch the particular student by simply passing the roll_number + 1 in the index
    # Below code works fast (if rollnumbers are sequential and in the format 1,2,3)
    # for row in sheet[roll_number+1]:
    #     # print(row[0].value)
    #     # print([x.value for x in row])
    #     print(row.value)

    # Fetch the row Linearly if roll numbers are of different format e.g. 21105002, CS001, etc
    for row in sheet:
        if row[0].value == roll_number:
            row_data = [x.value for x in row] # convert row into value list
            print(row_data[1])
            count=0
            for x in row_data:
                if x == 'P':
                    count+=1

    total_lecture = sheet.max_column - 2
    percentage = ceil((count/total_lecture)*100)
    detained_status = 'yes' if percentage < 75 else 'no'
    wb.close()
    return {'cid':course_code, 'cname':course_name,'stot_lect': count,'total_lect':total_lecture,'sper':percentage, 'detained_status':detained_status}