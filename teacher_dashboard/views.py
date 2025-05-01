from django.shortcuts import render, redirect
from django.http import HttpResponse
from django.contrib import messages
from django.contrib.auth.models import User, auth
from math import ceil
from openpyxl import load_workbook
from user_authentication.models import Stud   
# import pandas as pd
from openpyxl.utils import get_column_letter
import datetime as dt
from user_authentication.models import Record
# Create your views here.

excel_sheet_path='Excel_Sheets/main/'

def dashboard(request):
    if not request.user.is_authenticated:
        return redirect('login')
    
    teacher = Record.objects.filter(email = request.user.username).first()
    username = teacher.name
    course_name = teacher.course
    course_code = teacher.course_code
    sheet_name = course_code+'.xlsx'

    # Managing Sessions
    request.session['course_code'] = course_code
    request.session['course_name'] = course_name
    request.session['sheet_name'] = sheet_name
    try:
        wb = load_workbook(excel_sheet_path+sheet_name)
    except Exception as e:
        return HttpResponse("Excel sheet doesnt exist ! <br/> \
        It seems you have not conducted any lectures Please conduct atleast one lecture\
        <br/><a href='/logout'>Logout</a>")
    sheet = wb.active
    tot_lect = []
    record = []    
    tot_stud_on_a_day = []
    detained_student_mail = []
    total_lecture = sheet.max_column - 2 # reduce 2 to remove the 1st two column count
    #retrieving data from db , calc tot_lect , display on html
    i=2
    for u in Stud.objects.all():
        data = {}
        #info of each student as a key of dict
        data['sid'] = u.id
        data['sname'] = u.name
        #tot lect of each student as a key of dict
        count=0
        for row in sheet[i]:
            if row.value == "P":
                count+=1
        data['stot_lect'] = count
        #percentage claculation
        percentage = ceil((count/total_lecture)*100)
        data['sper'] = percentage
        detained_status=''
        if percentage < 75:
            detained_status='yes'
            detained_student_mail.append({'mail':u.email,'percentage':percentage})
        #appending the count to list.
        else:
            detained_status='no'
        data['detained_status']=detained_status
        tot_lect.append(count)
        i+=1
        record.append(data)
    # print(*record,sep='\n')
    request.session['total_lectures_attended'] = tot_lect
    # print(detained_student_mail)
    request.session['detained_student_mail'] = detained_student_mail
    for col in sheet.iter_cols(min_row=2,
                           max_row=sheet.max_row,
                           min_col=3,
                           max_col=sheet.max_column):
        cou = 0
        for data in col:
            if data.value == 'P':
                # print(data.value,end="")
                cou+=1        
        tot_stud_on_a_day.append(cou)
    # print(tot_stud_on_a_day)
    
    # record.append({'sid':1, 'sname':'sarthak','stot_lect': 2,'sper':'78'})
    
    total_students = sheet.max_row - 1
    today_date = dt.datetime.now()
    today_date = f'{today_date:%d-%m-%Y}'
    wb.close()
    return render(request,'teacher_dashboard/index.html',
    {
    'record':record,
    'username':username,
    'total_lecture':total_lecture,
    'total_students':total_students,
    'course_code':course_code,
    'course_name':course_name,
    'today_date':today_date,
    'graph_data':tot_stud_on_a_day
    })
    
def download(request):
    #download the excel file and append the tot lect col
    sheet_name = request.session['sheet_name']
    wb = load_workbook(excel_sheet_path+sheet_name)
    sheet = wb.active
    new_column = get_column_letter(sheet.max_column + 1)
    header_cell = new_column + '1'
    sheet[header_cell] = 'Total Lecture' 
    tot_lect = request.session['total_lectures_attended']  
    print('Total lecture',tot_lect)
    for i, value in enumerate(tot_lect, start=2):
        cell = new_column + str(i)
        sheet[cell] = value
        print(value)  
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=Student_Data.xlsx'
    wb.save(response)
    wb.close()
    # wb.save('Excel_Sheets/temp/Student_Data.xlsx') This downloads the file in the server so no need 
    return response

from django.core.mail import send_mass_mail #,EmailMessage
from django.conf import settings
def send_mail(request):
    subject='Detention Warning !'
    # message='Dear Student, Your attendance is low. Kindly communicate with respective teacher'
    from_email = settings.EMAIL_HOST_USER
    mail_data=[]
    c_code=request.session['course_code']
    c_name=request.session['course_name']
    mails = request.session['detained_student_mail']
    for x in mails:
        mail_data.append((
            subject,
            f'Dear Student, Your attendance is {x["percentage"]}%. Course:{c_code} - {c_name}. Kindly communicate with respective teacher',
            from_email,
            [x['mail']],
        ))
    print(mail_data)
    # send_mass_mail(tuple(mail_data),fail_silently=True)
    
    # mail = EmailMessage(subject, message, from_email, recipient_list)
    # mail.send()
    return redirect('teacher_dashboard')
    # return HttpResponse("Mail sent Successfully")