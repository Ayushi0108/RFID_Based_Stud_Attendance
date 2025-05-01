import os
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'RFID_Attendance_System.settings')
import django
django.setup()
from openpyxl import load_workbook
from user_authentication.models import Stud, Record
from math import ceil
from django.core.mail import send_mass_mail ,EmailMessage
from django.conf import settings
import schedule
import time

excel_sheet_path='Excel_Sheets/main/'

def send_email():
    course_codes=Record.objects.all().order_by('course_code').values('course_code','course')
    for course in course_codes:
        sheet_name = course['course_code']+'.xlsx'
        print(sheet_name)
        try:
            wb = load_workbook(excel_sheet_path+sheet_name)
        except Exception:            
            continue
        sheet = wb.active

        record=[]
        tot_lect=[]
        tot_stud_on_a_day = []
        detained_student_mail=[]

        total_lecture = sheet.max_column - 2
        i=2
        for u in Stud.objects.all():
            data = {}
                #info of each student as a key of dict
            data['sid'] = u.id
            data['sname'] = u.name
                #tot lect of each student as a key of dict
            count=0
            for row in sheet[i]:
                if row.value == "P":
                    count+=1
            data['stot_lect'] = count
                #percentage claculation
            percentage = ceil((count/total_lecture)*100)
            data['sper'] = percentage
            if percentage < 75:
                detained_student_mail.append({'mail':u.email,'percentage':percentage,})#'c_code':course['course_code'],'c_name':course['course']})
                #appending the count to list.
            tot_lect.append(count)
            i+=1
            record.append(data)
        # print(detained_student_mail)
        for student in detained_student_mail:
            send_email_to_student(student['mail'], student['percentage'], course['course_code'],course['course'])


def send_email_to_student(recipient, percentage,c_code,c_name):
    subject = 'Detention Warning !'
    message = f'Dear Student, Your attendance is {percentage}%. Course:{c_code} - {c_name}. Kindly communicate with respective teacher'
    from_email = settings.EMAIL_HOST_USER
    mail_data = (
        (subject, message, from_email, [recipient]),
    )
    print(mail_data)
    # send_mass_mail((mail_data,), fail_silently=True)


# schedule.every().day.at("08:00").do(send_email)

schedule.every(5).seconds.do(send_email)

# Schedule for partifular date also and the date will be set in database
# Or just hard code in the code :)

while True:
    schedule.run_pending()
    time.sleep(1)