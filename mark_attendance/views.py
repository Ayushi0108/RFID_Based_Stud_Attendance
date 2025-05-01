from django.shortcuts import render
from django.http import HttpResponse
from django.contrib import messages
# from django.contrib.auth.models import User, auth
from user_authentication.models import Record, Stud
#Current IP address;- 192.168.112.55:8000
# http://127.0.0.1:8000/mark/?rfid=d966514
def mark(request): #http://127.0.0.1:8000/mark?rfid=1234  #rf_10 -teacher ,  rf_1 , rf_2, rf_3 for student
    if request.method == 'GET':
        # Fetch rfid_number from URL
        rfid_number = request.GET.get('rfid')
        if rfid_number is None:
            rfid_number = "-1" # Set -1 if RFID not passed in URL
        teacher = Record.objects.filter(rfid=rfid_number).first()
        if teacher: # Is teacher
            print(teacher) # save the student attendance    
            roll_list = load_file()
            print(roll_list)
            sheet_name=teacher.course_code+'.xlsx'
            print(sheet_name)
            # from .save_excel_module import save
            # save(roll_list, sheet_name) # save the student data into excel sheet
            save_excel(roll_list, sheet_name)
            return HttpResponse("1")
        
        student = Stud.objects.filter(rfid=rfid_number).first()
        if student: # Is student
            print(student)
            save_roll(str(student.roll))
            return HttpResponse("1")

        # return HttpResponse("WHo are you ? <br/><a href='/logout'>Logout</a><br/>")
        return HttpResponse("0")

    return HttpResponse('Error no get request')
    # return HttpResponse("Attendance marked" + rfid_number + 
    # (student.name if student else (teacher.name if teacher else "No one") ))

def save_roll(roll: str):
    roll_list:list = load_file()    
    if roll in roll_list:
        roll_list.remove(roll)#comment this line if u dont want to remove the roll number if it is already present in the file
        pass
    else:
        roll_list.append(roll)
    with open('temp.txt','w') as my_file:        
        my_file.write(','.join(roll_list))

def load_file():
    roll_list: list=[]
    try:   
        with open('temp.txt','r') as my_file:
            data = my_file.readline()
            roll_list = data.split(',')
            if('') in roll_list:
                roll_list.remove('')
    except Exception as e:
        with open('temp.txt','w') as temp:
            pass
    return roll_list

def test(request):
    teachers = Record.objects.all().values('rfid','name')
    students = Stud.objects.all().values('rfid','name')
    return render(request, 'mark_attendance/test.html',
    {
        "teachers":teachers,
        "students":students
    })
    if request.method == 'GET':
        try:
            rfid_number = request.GET['rfid']
            #can also use request.GET.get() functioni it returns None if parameter not present
            #also no need to put that in try block
        except Exception as e:
            rfid_number = "-1"
    return HttpResponse("Test value passed : "+rfid_number)

import openpyxl
import datetime as dt
def save_excel(student_id, sheet_name):
    print(student_id, sheet_name)
    excel_sheet_path='Excel_Sheets/main/'
    # student_id = [1, 2, 3, 4, 7, 9, 10] # <- replace it here with the data received by the server
    date = dt.datetime.now()
    date = f"{date:%d-%m-%Y}"
    try:
        workbook = openpyxl.load_workbook(excel_sheet_path+sheet_name)
    except Exception:
        create_workbook(excel_sheet_path+sheet_name)
        workbook = openpyxl.load_workbook(excel_sheet_path+sheet_name)
    sheet = workbook.active
    val = sheet.max_column + 1
    sheet.cell(1, val).value = date           
    for row in range(2, sheet.max_row + 1):
        key = int(sheet.cell(row, 1).value)
        # print(i)
        if str(key) in student_id:
            print(key)
            sheet.cell(row, val).value = "P"
        else:
            sheet.cell(row, val).value = "A"
    workbook.save(excel_sheet_path+sheet_name)
    workbook.close()


from user_authentication.models import Record, Stud

def create_workbook(file):
    print(file)
    # file = 'Excel_Sheets/main/test_File.xlsx'
    workbook = openpyxl.Workbook()
    students_info = Stud.objects.all().order_by('roll').values('roll','name')
    sheet = workbook.worksheets[0]
    sheet.cell(row=1,column=1).value = 'Roll No'
    sheet.cell(row=1,column=2).value = 'Name'
    r_id=2
    for student in students_info:
        sheet.cell(row=r_id, column=1).value = student['roll']
        sheet.cell(row=r_id, column=2).value = student['name']
        r_id+=1

    workbook.save(file)
    workbook.close()